"""Spreadsheet helpers: the single output format and cell-level accuracy (no heavy imports).

Sheets are exchanged with the model as CSV text. Use xlsx_to_csv / csv_to_xlsx
(needs openpyxl) to convert real Excel files.
"""
import csv
import io

# The one output format every answer must follow. Edit to match your real sheet.
OUTPUT_COLUMNS = ["Date", "Vendor", "Description", "Quantity", "Unit Price", "Total"]

NUMERIC_TOLERANCE = 0.01


def parse_csv(text: str) -> list[list[str]]:
    rows = list(csv.reader(io.StringIO(text.strip("\n"))))
    return [r for r in rows if any(c.strip() for c in r)]


def _norm(cell: str) -> str:
    return " ".join(cell.split()).casefold()


def _as_number(cell: str) -> float | None:
    s = cell.strip().replace("$", "").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def cells_equal(pred: str, gold: str) -> bool:
    p, g = _as_number(pred), _as_number(gold)
    if p is not None and g is not None:
        return abs(p - g) <= NUMERIC_TOLERANCE
    return _norm(pred) == _norm(gold)


def header_matches(rows: list[list[str]]) -> bool:
    return bool(rows) and [_norm(c) for c in rows[0]] == [_norm(c) for c in OUTPUT_COLUMNS]


def cell_accuracy(pred_csv: str, gold_csv: str) -> float:
    """Fraction (0..1) of the gold sheet's cells the prediction got right.

    Compared by column name and row position over the OUTPUT_COLUMNS. Missing or
    extra rows count as wrong, so the denominator is max(gold rows, predicted rows).
    """
    pred, gold = parse_csv(pred_csv), parse_csv(gold_csv)
    if not pred or not gold:
        return 0.0
    pred_idx = {_norm(c): i for i, c in enumerate(pred[0])}
    gold_idx = {_norm(c): i for i, c in enumerate(gold[0])}
    pred_body, gold_body = pred[1:], gold[1:]

    n_rows = max(len(pred_body), len(gold_body))
    total = n_rows * len(OUTPUT_COLUMNS)
    if total == 0:
        return 1.0 if not pred_body and not gold_body else 0.0

    def cell(row, idx, col):
        i = idx.get(_norm(col))
        return row[i] if i is not None and i < len(row) else None

    correct = 0
    for r in range(min(len(pred_body), len(gold_body))):
        for col in OUTPUT_COLUMNS:
            p, g = cell(pred_body[r], pred_idx, col), cell(gold_body[r], gold_idx, col)
            if p is not None and g is not None and cells_equal(p, g):
                correct += 1
    return correct / total


def xlsx_to_csv(path: str, sheet: str | None = None) -> str:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)[sheet] if sheet else load_workbook(path, data_only=True).active
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    for row in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() for c in row):
            writer.writerow(["" if c is None else c for c in row])
    return buf.getvalue()


def csv_to_xlsx(csv_text: str, path: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in parse_csv(csv_text):
        ws.append(row)
    wb.save(path)
